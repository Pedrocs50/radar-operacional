"""Conteúdo do e-mail: visão geral do estoque, tabela em imagem, Excel e HTML.

O e-mail mostra TODAS as CDs abertas (mensagem["estoque"]) e destaca em vermelho as que acabaram
de chegar (mensagem["itens"]). Tudo aqui contém dado operacional em claro: os arquivos gerados
ficam na pasta temporária das prévias (retenção curta, apagada por `radar limpar`).
"""
from __future__ import annotations

import re
from collections import Counter
from datetime import date, datetime, time
from decimal import Decimal
from pathlib import Path

from jinja2 import Environment, FileSystemLoader, select_autoescape

from . import grafico, imagem, imagem_tabela
from ..engine import paineis
from ..engine.sinais import ATENCAO

TEMPLATES = Path(__file__).resolve().parent.parent / "templates"
FAIXAS = ["Hoje", "1 a 3 dias", "4 a 7 dias", "8 a 15 dias", "16 a 30 dias", "Mais de 30 dias"]
CAMPOS_PADRAO = {"dias": "Dias_em_Aberto", "faixa": "Faixa_de_Atraso",
                 "modelo": "Modelo_Aeronave", "ct_atual": "Descricao_CT_Atual"}
_ILEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def _cfg(m: dict) -> dict:
    return {**CAMPOS_PADRAO, **m.get("email", {}).get("campos", {})}


def _int(v):
    try:
        return int(v)
    except (TypeError, ValueError):
        return None


def chaves_novas(m: dict) -> set[str]:
    chave = m["campo_chave"]
    return {str(i.get(chave)) for i in m["itens"]}


def visao_geral(estoque: list[dict], c: dict, chave: str) -> dict:
    com_dias = [(d, l) for l in estoque if (d := _int(l.get(c["dias"]))) is not None]
    mais_antiga = max(com_dias, key=lambda x: x[0]) if com_dias else None
    faixas = Counter(str(l.get(c["faixa"]) or "-") for l in estoque)
    return {
        "total": len(estoque),
        "mais_antiga_dias": mais_antiga[0] if mais_antiga else None,
        "mais_antiga_nota": mais_antiga[1].get(chave) if mais_antiga else None,
        "media_dias": round(sum(d for d, _ in com_dias) / len(com_dias), 1) if com_dias else None,
        "faixas": [(f, faixas[f]) for f in FAIXAS if faixas.get(f)],
        "modelos": Counter(str(l.get(c["modelo"]) or "-") for l in estoque).most_common(),
        "cts": Counter(str(l.get(c["ct_atual"]) or "-") for l in estoque).most_common(5),
    }


def _celula(v):
    if v is None or isinstance(v, (int, float, bool, date, datetime, time)):
        return v
    if isinstance(v, Decimal):
        return float(v)
    return _ILEGAL.sub("", str(v))


def _com_atencao(m: dict) -> list[dict]:
    """Estoque com a coluna "Atenção" (tags dos sinais), se a tarefa configura sinais."""
    marcas = m.get("marcas")
    if marcas is None:
        return m["estoque"]
    return [{**l, ATENCAO: ", ".join(marcas.get(str(l.get(m["campo_chave"])), []))} for l in m["estoque"]]


def gerar_xlsx(m: dict, destino: Path) -> Path:
    """Excel com TODAS as colunas do SQL, uma linha por CD; as novas em vermelho."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    estoque, novas, chave = _com_atencao(m), chaves_novas(m), m["campo_chave"]
    colunas = list(estoque[0].keys()) if estoque else list(m["colunas"])
    if ATENCAO in colunas:  # a marca de atenção vem logo depois da situação
        colunas.remove(ATENCAO)
        colunas.insert(0, ATENCAO)
    wb = Workbook()
    ws = wb.active
    ws.title = "CDs abertas"
    ws.append(["Situação"] + [m.get("rotulos", {}).get(c, c) for c in colunas])
    for cel in ws[1]:
        cel.font = Font(bold=True, color="FFFFFF")
        cel.fill = PatternFill("solid", fgColor="1F3A5F")
    vermelho = PatternFill("solid", fgColor="FDE2E2")
    amarelo = PatternFill("solid", fgColor="FFF3CD")
    for l in estoque:
        nova = str(l.get(chave)) in novas
        ws.append(["NOVA" if nova else ""] + [_celula(l.get(c)) for c in colunas])
        if nova:
            for cel in ws[ws.max_row]:
                cel.fill = vermelho
                cel.font = Font(bold=True, color="9B1C1C")
        elif l.get(ATENCAO):
            for cel in ws[ws.max_row]:
                cel.fill = amarelo
    for i, col in enumerate(ws.columns, start=1):
        largura = max((len(str(c.value)) for c in col if c.value is not None), default=8)
        ws.column_dimensions[get_column_letter(i)].width = min(largura + 2, 50)
        for c in col:
            c.alignment = Alignment(vertical="top")
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    destino.parent.mkdir(parents=True, exist_ok=True)
    wb.save(destino)
    return destino


def renderizar_html(m: dict, img_src: str, grafico_src: str | None = None, painel_src: str | None = None) -> str:
    c = _cfg(m)
    env = Environment(loader=FileSystemLoader(str(TEMPLATES)), autoescape=select_autoescape(["html", "j2"]))
    rotulos = m.get("rotulos", {})
    return env.get_template("email.html.j2").render(
        m=m, geral=visao_geral(m["estoque"], c, m["campo_chave"]), img_src=img_src, grafico_src=grafico_src, painel_src=painel_src,
        painel_alt=paineis.texto_alternativo(m["paineis"]) if m.get("paineis") else "",
        detalhe=[(rotulos.get(col, col), col) for col in m["colunas"]],
    )


def assunto(m: dict) -> str:
    teste = " · TESTE" if m.get("sandbox") else ""
    if m.get("tipo") == "resumo":
        g = m["geral"]
        antiga = f" · mais antiga {g['mais_antiga_dias']} dias" if g.get("mais_antiga_dias") is not None else ""
        return f"[Radar{teste}] {m['titulo']}: {len(m['estoque'])} aberta(s){antiga}"
    return f"[Radar{teste}] {m['titulo']}: {len(m['itens'])} nova(s) · {len(m['estoque'])} aberta(s)"


def gerar_arquivos(m: dict, pasta: Path, nome_base: str) -> dict:
    """Gera a tabela em PNG e o Excel. Retorna os caminhos."""
    pasta = Path(pasta)
    colunas = list(m.get("email", {}).get("colunas_tabela") or m["colunas"])
    if m.get("marcas") is not None:
        colunas.insert(1, ATENCAO)
    png = imagem_tabela.gerar_tabela(_com_atencao(m), colunas, pasta / f"{nome_base}_tabela.png", m.get("rotulos"),
                              chaves_novas(m), m["campo_chave"])
    xlsx = gerar_xlsx(m, pasta / f"{nome_base}_cds_abertas.xlsx")
    arquivos = {"png": png, "xlsx": xlsx}
    if m.get("paineis"):
        arquivos["painel"] = grafico.gerar_paineis(m["paineis"], pasta / f"{nome_base}_painel.png")
    ind = m.get("indicadores")
    if ind and ind.get("tem_dados"):
        arquivos["grafico"] = grafico.gerar(ind["semanas"], pasta / f"{nome_base}_grafico.png")
    return arquivos


def gerar_previa(m: dict, pasta: Path, nome_base: str) -> Path:
    """Prévia local do e-mail (HTML + PNG + Excel), sem enviar nada."""
    arq = gerar_arquivos(m, pasta, nome_base)
    destino = Path(pasta) / f"{nome_base}_email.html"
    destino.write_text(renderizar_html(m, arq["png"].name, arq["grafico"].name if "grafico" in arq else None,
                                       arq["painel"].name if "painel" in arq else None), encoding="utf-8")
    return destino
