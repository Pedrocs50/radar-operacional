import pytest

from radar.channels import email_conteudo as ec
from radar.channels import email_outlook
from radar.channels.base import ErroDeEnvio
from radar.engine.runner import executar
from tests.conftest import FonteFalsa


def linha(nota, dias, faixa, modelo="EJETS E2", ct="CT EXEMPLO", problema="problema"):
    return {"Nota_QM": nota, "Dias_em_Aberto": dias, "Faixa_de_Atraso": faixa, "Modelo_Aeronave": modelo,
            "Descricao_CT_Atual": ct, "Descricao_do_Problema": problema, "Fluxo": None}


ESTOQUE = [linha("1", 40, "Mais de 30 dias"), linha("2", 5, "4 a 7 dias", "EJETS E1"), linha("3", 0, "Hoje")]


def mensagem(novas=("3",)):
    return {
        "titulo": "Nova CD", "marca": "RADAR OPERACIONAL", "dados_de": "07/10/2026 13:00", "sandbox": True,
        "colunas": ["Nota_QM", "Descricao_do_Problema"], "rotulos": {"Nota_QM": "Nota QM"},
        "estoque": ESTOQUE, "itens": [l for l in ESTOQUE if l["Nota_QM"] in novas], "campo_chave": "Nota_QM",
        "email": {"colunas_tabela": ["Nota_QM", "Dias_em_Aberto", "Descricao_do_Problema"]},
    }


def test_visao_geral():
    g = ec.visao_geral(ESTOQUE, ec.CAMPOS_PADRAO, "Nota_QM")
    assert g["total"] == 3 and g["mais_antiga_dias"] == 40 and g["mais_antiga_nota"] == "1"
    assert g["media_dias"] == 15.0
    assert dict(g["modelos"]) == {"EJETS E2": 2, "EJETS E1": 1}
    assert ("Hoje", 1) in g["faixas"]


def test_excel_marca_a_nova_em_vermelho_e_traz_todas_as_colunas(tmp_path):
    from openpyxl import load_workbook

    ec.gerar_xlsx(mensagem(), tmp_path / "x.xlsx")
    ws = load_workbook(tmp_path / "x.xlsx").active
    assert ws.max_row == 4 and ws.max_column == 1 + len(ESTOQUE[0])
    assert ws["A4"].value == "NOVA" and ws["A2"].value in (None, "")
    assert ws["B4"].fill.fgColor.rgb.endswith("FDE2E2")


def test_html_escapa_e_mostra_visao_geral():
    m = mensagem()
    m["itens"][0]["Descricao_do_Problema"] = "<script>x</script>"
    html = ec.renderizar_html(m, "cid:tabela")
    assert "<script>x</script>" not in html and "Visão geral" in html
    assert "40 dias" in html and 'src="cid:tabela"' in html


def test_html_usa_aging_e_layout_compacto():
    html = ec.renderizar_html(mensagem(), "cid:tabela")
    assert "Aging" in html and "Dias em aberto" not in html
    assert "Chegou agora: 1 CD(s) nova(s)" in html
    assert "Onde estão (CT atual):" in html
    assert "<li>Hoje: <b>1</b></li>" in html or "Hoje: <b>1</b>" in html


def test_arquivos_e_previa(tmp_path):
    arq = ec.gerar_arquivos(mensagem(), tmp_path, "t")
    assert arq["png"].read_bytes()[:4] == b"\x89PNG" and arq["xlsx"].exists()
    assert ec.gerar_previa(mensagem(), tmp_path, "t").name == "t_email.html"


def test_enviar_monta_o_pedido_e_nao_envia_de_verdade(tmp_path):
    capturado = {}
    email_outlook.enviar(mensagem(), "a@x.com; b@y.com", tmp_path, "t", enviador=lambda d, p, n: capturado.update(d))
    assert capturado["para"] == ["a@x.com", "b@y.com"]
    assert capturado["assunto"] == "[Radar · TESTE] Nova CD: 1 nova(s) · 3 aberta(s)"
    assert capturado["inline"][0]["cid"] == "tabela" and capturado["anexos"][0].endswith(".xlsx")


def test_destinatario_invalido(tmp_path):
    with pytest.raises(ErroDeEnvio) as e:
        email_outlook.enviar(mensagem(), "isto nao e email", tmp_path, "t", enviador=lambda *a: None)
    assert e.value.codigo == "EMAIL_DESTINATARIO_INVALIDO"


def test_runner_com_canal_email(task, storage, crypto, settings, tmp_path):
    task = {**task, "canais": ["preview", "email:teste"],
            "mensagem": {**task["mensagem"], "campos_exibidos": ["Nota", "Descricao_do_Problema"]}}
    settings = {**settings, "sandbox": {"ativo": True, "canais_permitidos": ["preview", "email:teste"]}}
    linhas = [{"Nota": str(n), "Descricao_do_Problema": "p", "Data_da_CD": "x"} for n in (1, 2, 3)]
    recebido = []
    r = executar(task, FonteFalsa(linhas), storage, crypto, settings, dry_run=False,
                 enviadores={"email:teste": recebido.append}, limite=1)
    assert r["enviados"] == ["email:teste"] and r["email_previa"]
    assert len(recebido[0]["estoque"]) == 3 and len(recebido[0]["itens"]) == 1  # estoque inteiro, 1 nova
